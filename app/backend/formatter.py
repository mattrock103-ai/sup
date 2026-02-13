from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph

SECTION_KEYS = [
    "SKILLS",
    "EXPERIENCE",
    "EDUCATION",
    "QUALIFICATIONS",
    "CERTIFICATIONS",
    "ADDITIONAL_INFO",
]

SIMPLE_KEYS = ["NAME", "TITLE", "LOCATION", "CONTACT_LINE", "SUMMARY"]


def _iter_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _insert_paragraph_after(paragraph: Paragraph, text: str = "", style_name: str | None = None) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.clear()
    if text:
        new_para.add_run(text)
    if style_name:
        try:
            new_para.style = style_name
        except Exception:
            pass
    return new_para


def _find_style_name(doc: Document, preferred: list[str], fallback: str | None = None) -> str | None:
    style_names = {s.name for s in doc.styles}
    for name in preferred:
        if name in style_names:
            return name
    return fallback


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        out = []
        for item in value:
            if isinstance(item, str):
                out.append(item)
            else:
                out.append(str(item))
        return out
    if isinstance(value, str):
        return [line.strip() for line in value.splitlines() if line.strip()]
    return [str(value)]


def _replace_simple_placeholders(doc: Document, data: dict[str, Any]) -> None:
    for p in _iter_paragraphs(doc):
        original = p.text
        replaced = original
        for key in SIMPLE_KEYS:
            token = f"{{{{{key}}}}}"
            value = str(data.get(key.lower(), "") or "")
            replaced = replaced.replace(token, value)
        if replaced != original:
            p.clear()
            p.add_run(replaced)


def _render_structured_block(paragraph: Paragraph, key: str, data: dict[str, Any], doc: Document) -> None:
    key_lower = key.lower()

    body_style = paragraph.style.name if paragraph.style else None
    bullet_style = _find_style_name(doc, ["List Bullet", "List Paragraph"], fallback=body_style)

    paragraph.clear()

    if key == "SKILLS":
        for idx, skill in enumerate(_as_list(data.get("skills"))):
            target = paragraph if idx == 0 else _insert_paragraph_after(paragraph, style_name=bullet_style)
            target.add_run(skill)
            target.style = bullet_style
            paragraph = target
        return

    if key == "EXPERIENCE":
        exp_items = data.get("experience") or []
        if isinstance(exp_items, str):
            exp_items = [{"title": "Experience", "company": "", "dates": "", "bullets": [exp_items]}]

        first = True
        current_anchor = paragraph
        for role in exp_items:
            title_line = " | ".join([part for part in [role.get("title", ""), role.get("company", ""), role.get("dates", "")] if part]).strip(" |")
            if not title_line:
                title_line = "Experience"

            role_para = current_anchor if first else _insert_paragraph_after(current_anchor, style_name=body_style)
            role_para.add_run(title_line)
            role_para.style = body_style
            current_anchor = role_para
            first = False

            for bullet in role.get("bullets", []):
                bullet_para = _insert_paragraph_after(current_anchor, text=str(bullet), style_name=bullet_style)
                bullet_para.style = bullet_style
                current_anchor = bullet_para
        return

    if key == "EDUCATION":
        education = data.get("education") or []
        if isinstance(education, str):
            education = [{"school": education, "degree": "", "dates": ""}]
        for idx, item in enumerate(education):
            line = " | ".join([part for part in [item.get("school", ""), item.get("degree", ""), item.get("dates", "")] if part])
            target = paragraph if idx == 0 else _insert_paragraph_after(paragraph, style_name=body_style)
            target.add_run(line)
            target.style = body_style
            paragraph = target
        return

    raw_values = _as_list(data.get(key_lower, ""))
    for idx, val in enumerate(raw_values):
        style = bullet_style if key in {"QUALIFICATIONS", "CERTIFICATIONS"} else body_style
        target = paragraph if idx == 0 else _insert_paragraph_after(paragraph, style_name=style)
        target.add_run(val)
        if style:
            target.style = style
        paragraph = target


def render_cv(template_path: Path, parsed_data: dict[str, Any], output_dir: Path) -> Path:
    doc = Document(str(template_path))

    _replace_simple_placeholders(doc, parsed_data)

    for p in list(_iter_paragraphs(doc)):
        text = p.text.strip()
        for key in SECTION_KEYS:
            token = f"{{{{{key}}}}}"
            if token in text:
                _render_structured_block(p, key, parsed_data, doc)

    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
    output_path = output_dir / f"formatted_cv_{timestamp}.docx"
    doc.save(str(output_path))
    return output_path


def build_default_template_if_missing(path: Path) -> None:
    if path.exists():
        return

    doc = Document()
    doc.add_heading("{{NAME}}", level=0)
    doc.add_paragraph("{{TITLE}}")
    doc.add_paragraph("{{LOCATION}}")
    doc.add_paragraph("{{CONTACT_LINE}}")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph("{{SUMMARY}}")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph("{{SKILLS}}", style="List Bullet")

    doc.add_heading("Experience", level=1)
    doc.add_paragraph("{{EXPERIENCE}}")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("{{EDUCATION}}")

    doc.add_heading("Qualifications", level=1)
    doc.add_paragraph("{{QUALIFICATIONS}}", style="List Bullet")

    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("{{CERTIFICATIONS}}", style="List Bullet")

    doc.add_heading("Additional Information", level=1)
    doc.add_paragraph("{{ADDITIONAL_INFO}}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
